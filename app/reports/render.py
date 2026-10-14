"""PDF (WeasyPrint) and Excel (openpyxl) output from the same report data."""

from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import brand
from app.reports import charts
from app.reports.data import BUNDLE_TITLES, human

_env = Environment(loader=FileSystemLoader(str(Path(__file__).resolve().parent.parent / "templates" / "reports")),
                   autoescape=select_autoescape(["html"]))


def _local(dt: datetime | None, tz: str, fmt: str = "%Y-%m-%d %H:%M:%S") -> str:
    return dt.astimezone(ZoneInfo(tz)).strftime(fmt) if dt else ""


def _pct(v: float | None) -> str:
    return "" if v is None else f"{v * 100:.3f} %"


_env.filters["local"] = _local
_env.filters["human"] = human
_env.filters["pct"] = _pct


def _gb(n) -> str:
    return f"{(n or 0) / 1e9:,.2f} GB"


def daily_svg(days: list[dict], width: int = 640, height: int = 170) -> str:
    """Daily download/upload bars as inline SVG (WeasyPrint renders it; no JavaScript in PDFs)."""
    if not days:
        return ""
    top = max(max(d["rx"], d["tx"]) for d in days) or 1
    left, bottom, pad = 52, 22, 2
    slot = (width - left) / len(days)
    bw = max(1.0, slot / 2 - pad)
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" font-family="Montserrat, DejaVu Sans" font-size="8">']
    for frac in (0, 0.5, 1):
        y = (height - bottom) * (1 - frac) + 4
        parts.append(f'<line x1="{left}" x2="{width}" y1="{y:.1f}" y2="{y:.1f}" stroke="#e1e0d9" stroke-width="0.5"/>')
        parts.append(f'<text x="{left - 4}" y="{y + 3:.1f}" text-anchor="end" fill="#898781">{top * frac / 1e9:.1f} GB</text>')
    for i, d in enumerate(days):
        x = left + i * slot
        for j, (key, color) in enumerate((("rx", "#2a78d6"), ("tx", "#eb6834"))):
            h = (height - bottom - 4) * d[key] / top
            parts.append(f'<rect x="{x + j * (bw + pad):.1f}" y="{height - bottom - h:.1f}" width="{bw:.1f}" height="{h:.1f}" fill="{color}"/>')
        if len(days) <= 16 or i % max(1, len(days) // 10) == 0:
            parts.append(f'<text x="{x + slot / 2:.1f}" y="{height - 8}" text-anchor="middle" fill="#52514e">{d["date"]:%d %b}</text>')
    parts.append("</svg>")
    return "".join(parts)


def render_chart(b: dict, stacked: bool = False) -> str:
    kind = b["chart"]
    if kind == "donut":
        return charts.donut(b["items"], b.get("unit", ""), palette=b.get("palette"), stacked=stacked)
    if kind == "line":
        return charts.line_chart(b["x"], b["series"], b.get("unit", ""), area=b.get("area", False))
    if kind == "stacked":
        return charts.stacked_bars(b["x"], b["series"], b.get("unit", ""), palette=b.get("palette"))
    if kind == "hbars":
        return charts.hbars(b["items"], b.get("unit", ""), colour=b.get("colour", charts.SERIES[0]))
    return ""


def pair_pies(blocks: list[dict]) -> list[dict]:
    """Consecutive pie charts go two to a row ({"kind": "pies"}); a pie without a neighbour stands alone, centred."""
    out: list[dict] = []
    for b in blocks:
        if b.get("kind") == "chart" and b.get("chart") == "donut":
            if out and out[-1]["kind"] == "pies" and len(out[-1]["pies"]) < 2:
                out[-1]["pies"].append(b)
            else:
                out.append({"kind": "pies", "pies": [b]})
        else:
            out.append(b)
    return out


_env.filters["gb"] = _gb
_env.filters["pair_pies"] = pair_pies
_env.globals["daily_svg"] = daily_svg
_env.globals["render_chart"] = render_chart
_env.globals["bundle_titles"] = BUNDLE_TITLES
_env.globals["brand"] = brand
_env.globals["logo_uri"] = brand.asset("logo.png").as_uri()
_env.globals["cover_uri"] = brand.asset("cover.jpg").as_uri() if brand.asset("cover.jpg").exists() else None
_env.globals["one_second"] = timedelta(seconds=1)


def render_html(report: dict) -> str:
    return _env.get_template("report.html").render(r=report, tz=report["client"]["timezone"])


def write_pdf(report: dict, path: Path) -> None:
    from weasyprint import HTML  # imported lazily: needs Pango, present only in the app image

    path.parent.mkdir(parents=True, exist_ok=True)
    # base_url: the cover image, logo and Montserrat fonts in app/assets
    HTML(string=render_html(report), base_url=brand.ASSETS.as_uri() + "/").write_pdf(str(path))


# ---------------------------------------------------------------- Excel
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=brand.RED[1:])


def _safe(v):
    """Strings that look like formulas are stored as text, never evaluated."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, Alignment(vertical="center")
    for r in rows:
        ws.append([_safe(v) for v in r])
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.freeze_panes = "A2"


def write_xlsx(report: dict, path: Path) -> None:
    tz = report["client"]["timezone"]
    loc = lambda dt: dt.astimezone(ZoneInfo(tz)).replace(tzinfo=None) if dt else None  # noqa: E731
    wb = Workbook()
    info = wb.active
    info.title = "Report"
    for row in (("Client", report["client"]["name"]), ("Report", report["type"]),
                ("Period start", loc(report["start"])), ("Period end", loc(report["end"])),
                ("Time zone", tz), ("Generated", loc(report["generated"]))):
        info.append(row)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 16, 30
    used: set[str] = set()
    for s in report["sections"]:
        if s.get("generic"):
            _generic_sheets(wb, s, used)
            continue
        t = s["type"]
        if t == "internet_usage":
            _sheet(wb, "Internet usage", ["Gateway", "WAN", "Name", "Download (GB)", "Upload (GB)", "Avg down (Mbps)",
                                          "Avg up (Mbps)", "95th pct down (Mbps)", "95th pct up (Mbps)", "Peak down (Mbps)",
                                          "Peak up (Mbps)", "Packets in", "Packets out", "Drops", "Errors", "Availability %"],
                   [[w["device"], w["wan"], w["label"], round(w["rx"] / 1e9, 3), round(w["tx"] / 1e9, 3),
                     round(w["avg_rx_mbps"], 2), round(w["avg_tx_mbps"], 2), round(w["p95_rx_mbps"], 2), round(w["p95_tx_mbps"], 2),
                     round(w["peak_rx_mbps"], 2), round(w["peak_tx_mbps"], 2), w["prx"], w["ptx"], w["drops"], w["errors"],
                     round(w["availability"] * 100, 3) if w["availability"] is not None else None] for w in s["wans"]])
            _sheet(wb, "Internet daily", ["Gateway", "WAN", "Date", "Download (GB)", "Upload (GB)", "Peak down (Mbps)", "Peak up (Mbps)"],
                   [[w["device"], w["wan"], d["date"], round(d["rx"] / 1e9, 3), round(d["tx"] / 1e9, 3),
                     round(d["peak_rx"], 2), round(d["peak_tx"], 2)] for w in s["wans"] for d in w["days"]])
            _sheet(wb, "Internet outages", ["Gateway", "WAN", "Start", "End", "Duration (s)"],
                   [[w["device"], w["wan"], loc(a), loc(b), round((b - a).total_seconds())] for w in s["wans"] for a, b in w["outages"]])
        elif t == "device_uptime":
            _sheet(wb, "Device uptime", ["Device", "Model", "Firmware", "Availability %", "Up (s)", "Down (s)",
                                         "Maintenance (s)", "No data (s)", "Outages"],
                   [[d["name"], d["model"], d["firmware"], round(d["availability"] * 100, 3) if d["availability"] is not None else None,
                     round(d["up_s"]), round(d["down_s"]), round(d["maintenance_s"]), round(d["no_data_s"]), d["drops"]]
                    for d in s["devices"]])
            _sheet(wb, "Device outages", ["Device", "Start", "End", "Duration (s)", "Duration"],
                   [[d["name"], loc(a), loc(b), round((b - a).total_seconds()), human((b - a).total_seconds())]
                    for d in s["devices"] for a, b in d["outages"]])
        elif t == "vpn_uptime":
            _sheet(wb, "VPN uptime", ["Device", "Tunnel", "Peer", "Availability %", "Up (s)", "Down (s)", "No data (s)",
                                      "Maintenance (s)", "Drops"],
                   [[x["device"], x["tunnel"], x["peer"], round(x["availability"] * 100, 3) if x["availability"] is not None else None,
                     round(x["up_s"]), round(x["down_s"]), round(x["unknown_s"]), round(x["maintenance_s"]), x["drops"]]
                    for x in s["tunnels"]])
            _sheet(wb, "VPN downtime", ["Device", "Tunnel", "Start", "End", "Duration (s)"],
                   [[x["device"], x["tunnel"], loc(a), loc(b), round((b - a).total_seconds())]
                    for x in s["tunnels"] for a, b in x["outages"]])
        elif t == "vpn_sessions":
            _sheet(wb, "VPN sessions", ["User", "Source IP", "VPN address", "Device", "Start", "End", "Duration (s)", "Status", "Note"],
                   [[x["user_name"], x["source_ip"], x["assigned_ip"], x["device"], loc(x["started_at"]), loc(x["ended_at"]),
                     round(x["duration_s"]) if x["duration_s"] is not None else None, x["status"], x["close_reason"]]
                    for x in s["sessions"]])
            _sheet(wb, "VPN failed sources", ["Source IP", "Country", "Network", "Attempts"],
                   [[x["ip"], x["country"], x["org"], x["n"]] for x in s["failed_sources"]])
        elif t == "device_traffic":
            _sheet(wb, "Top sources", ["Device", "Source IP", "Country", "Events", "Blocked/failed", "Bytes"],
                   [[d["device"], x["ip"], x["country"], x["events"], x["blocked"], x["bytes"]] for d in s["devices"] for x in d["sources"]])
            _sheet(wb, "Top destinations", ["Device", "Destination IP", "Port", "Country", "Events", "Blocked/failed", "Bytes"],
                   [[d["device"], x["ip"], x["port"], x["country"], x["events"], x["blocked"], x["bytes"]]
                    for d in s["devices"] for x in d["destinations"]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


ABBR = {"executive_summary": "Exec", "traffic_bandwidth": "Traffic", "security_attacks": "Security", "rule_policy": "Rules",
        "vpn_remote_access": "VPN", "compliance_audit": "Compliance", "web_app_usage": "Web", "cctv_access": "CCTV", "ha_system": "System"}


def _sheet_name(title: str, used: set[str]) -> str:
    base = "".join(ch for ch in title if ch not in '[]:*?/\\')[:31].strip() or "Sheet"
    name, i = base, 2
    while name.lower() in used:
        suffix = f" {i}"
        name, i = base[:31 - len(suffix)] + suffix, i + 1
    used.add(name.lower())
    return name


def _generic_sheets(wb: Workbook, s: dict, used: set[str]) -> None:
    """KPIs and findings on one sheet; every table and every chart's data on its own sheet."""
    ab = ABBR.get(s["type"], s["type"][:8])
    rows = [["KPI", k["label"], k["value"], k["note"] or ""] for k in s["kpis"]]
    rows += [["Finding", f[0], f[1], ""] for f in s["findings"]]
    _sheet(wb, _sheet_name(f"{ab} summary", used), ["Type", "Item", "Value", "Note"], rows)
    for b in s["blocks"]:
        if b["kind"] == "table" and b["rows"]:
            _sheet(wb, _sheet_name(f"{ab} {b['title']}", used), b["headers"], b["rows"])
        elif b["kind"] == "chart":
            if b["chart"] in ("donut", "hbars"):
                _sheet(wb, _sheet_name(f"{ab} {b['title']}", used), ["Item", "Value"], [[k, v] for k, v in b["items"]])
            elif b["chart"] in ("line", "stacked") and b["series"]:
                _sheet(wb, _sheet_name(f"{ab} {b['title']}", used), ["Date"] + [n for n, _ in b["series"]],
                       [[x] + [vals[i] if i < len(vals) else None for _, vals in b["series"]] for i, x in enumerate(b["x"])])

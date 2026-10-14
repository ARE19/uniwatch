"""Clients / devices / unknown-sender pages."""

from datetime import timedelta

from sqlalchemy import text

from tests.conftest import csrf_from, full_login


async def _admin(client, make_user):
    await full_login(client, make_user(username="root-admin", role="admin"))
    return csrf_from((await client.get("/")).text)


async def test_admin_creates_client_and_device(client, make_user, owner_engine, audit_rows):
    token = await _admin(client, make_user)
    r = await client.post("/admin/clients", data={"name": "Northwind HK", "slug": "nwhk", "timezone": "Asia/Hong_Kong",
                                                   "retention_days": "90", "csrf_token": token})
    assert r.status_code == 303
    with owner_engine.connect() as c:
        cid = c.execute(text("SELECT id FROM clients WHERE slug='nwhk'")).scalar()
    r = await client.post("/sources/devices", data={"client_id": str(cid), "name": "NW-HK", "vendor": "unifi",
                                                    "model": "UDM-Pro-SE", "syslog_source_ip": "192.168.1.1",
                                                    "csrf_token": token})
    assert r.status_code == 303
    page = await client.get("/sources")
    assert "NW-HK" in page.text and "192.168.1.1" in page.text
    assert audit_rows("client_create") and audit_rows("device_create")


async def test_client_validation(client, make_user):
    token = await _admin(client, make_user)
    r = await client.post("/admin/clients", data={"name": "X", "slug": "Bad Slug!", "timezone": "Mars/Base",
                                                   "retention_days": "0", "csrf_token": token})
    assert r.status_code == 400
    assert "Slug" in r.text and "time zone" in r.text and "Retention" in r.text


async def test_duplicate_source_ip_rejected(client, make_user, owner_engine):
    token = await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug) VALUES ('A','aa') RETURNING id")).scalar()
        c.execute(text("INSERT INTO devices (client_id, name, vendor, syslog_source_ip) VALUES (:c,'one','unifi','10.0.0.1')"), {"c": cid})
    r = await client.post("/sources/devices", data={"client_id": str(cid), "name": "two", "vendor": "unifi",
                                                    "syslog_source_ip": "10.0.0.1", "csrf_token": token})
    assert r.status_code == 400 and "already the syslog source" in r.text
    r = await client.post("/sources/devices", data={"client_id": str(cid), "name": "three", "vendor": "unifi",
                                                    "syslog_source_ip": "not-an-ip", "csrf_token": token})
    assert r.status_code == 400 and "not a valid IP" in r.text


async def test_unknown_sender_listed_then_resolved_by_registering(client, make_user, owner_engine):
    token = await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug) VALUES ('A','aa') RETURNING id")).scalar()
        c.execute(text("INSERT INTO unknown_sources (source_ip, first_seen, last_seen, count, sample) "
                       "VALUES ('192.168.1.102', now(), now(), 7, '<13>NWHKWAP-01 hello')"))
    page = await client.get("/sources")
    assert "192.168.1.102" in page.text and "NWHKWAP-01 hello" in page.text
    assert "/sources/add?register=192.168.1.102" in page.text
    assert 'value="192.168.1.102"' in (await client.get("/sources/add?register=192.168.1.102")).text
    await client.post("/sources/devices", data={"client_id": str(cid), "name": "NWHKWAP-01", "vendor": "unifi",
                                                "syslog_source_ip": "192.168.1.102", "csrf_token": token})
    with owner_engine.connect() as c:
        assert c.execute(text("SELECT count(*) FROM unknown_sources")).scalar() == 0


async def test_ignore_unknown_sender(client, make_user, owner_engine):
    token = await _admin(client, make_user)
    with owner_engine.begin() as c:
        c.execute(text("INSERT INTO unknown_sources (source_ip, first_seen, last_seen, count, sample) "
                       "VALUES ('10.9.9.9', now(), now(), 1, 'noise')"))
    r = await client.post("/sources/unknown/ignore", data={"source_ip": "10.9.9.9", "csrf_token": token})
    assert r.status_code == 303
    assert "10.9.9.9" not in (await client.get("/sources")).text


async def test_raw_viewer_shows_lines_in_client_timezone(client, make_user, owner_engine):
    await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug, timezone) VALUES ('A','aa','Asia/Hong_Kong') RETURNING id")).scalar()
        did = c.execute(text("INSERT INTO devices (client_id, name, vendor) VALUES (:c,'udm','unifi') RETURNING id"), {"c": cid}).scalar()
        c.execute(text("INSERT INTO raw_logs (ts_received, client_id, device_id, source_ip, transport, raw) "
                       "VALUES ('2026-10-06T04:00:00Z', :c, :d, '192.168.1.1', 'udp', '<13>hello raw')"), {"c": cid, "d": did})
    r = await client.get(f"/sources/raw?device_id={did}")
    assert r.status_code == 200 and "&lt;13&gt;hello raw" in r.text
    assert "2026-10-06 12:00:00 HKT" in r.text  # 04:00Z shown as Hong Kong time


async def test_engineer_sees_sources_but_not_clients(client, make_user):
    await full_login(client, make_user(role="engineer"))
    assert (await client.get("/sources")).status_code == 200
    assert (await client.get("/admin/clients")).status_code == 403


async def test_client_readonly_cannot_see_sources(client, make_user):
    await full_login(client, make_user(role="client_readonly"))
    assert (await client.get("/sources")).status_code == 403
    assert (await client.get("/sources/raw?device_id=1")).status_code == 403
    assert (await client.get("/devices")).status_code == 403 and (await client.get("/sources/add")).status_code == 403


async def test_maintenance_window_create_and_delete(client, make_user, owner_engine, audit_rows):
    token = await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug, timezone) VALUES ('A','aa','Asia/Hong_Kong') RETURNING id")).scalar()
    r = await client.post("/sources/maintenance", data={"client_id": str(cid), "device_id": "", "starts": "2026-10-10T22:00",
                                                        "ends": "2026-10-10T23:30", "reason": "firmware upgrade", "csrf_token": token})
    assert r.status_code == 303
    with owner_engine.connect() as c:
        s, e = c.execute(text("SELECT starts_at, ends_at FROM maintenance_windows")).one()
    assert s.isoformat() == "2026-10-10T14:00:00+00:00" and e - s == timedelta(minutes=90)  # 22:00 HKT = 14:00Z
    page = await client.get("/sources/maintenance")
    assert "firmware upgrade" in page.text and "all devices" in page.text
    bad = await client.post("/sources/maintenance", data={"client_id": str(cid), "starts": "2026-10-10T23:00",
                                                          "ends": "2026-10-10T22:00", "csrf_token": token})
    assert bad.status_code == 400
    mw = page.text.split('action="/sources/maintenance/')[1].split("/")[0]
    assert (await client.post(f"/sources/maintenance/{mw}/delete", data={"csrf_token": token})).status_code == 303
    assert audit_rows("maintenance_create") and audit_rows("maintenance_delete")


async def test_discovered_devices_listed_with_register_link(client, make_user, owner_engine):
    await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug) VALUES ('A','aa') RETURNING id")).scalar()
        gw = c.execute(text("INSERT INTO devices (client_id, name, vendor) VALUES (:c,'gw','unifi') RETURNING id"), {"c": cid}).scalar()
        c.execute(text("INSERT INTO device_inventory (poll_device_id, mac, name, model, ip, connected, seen_at) "
                       "VALUES (:g, 'aa:bb:cc:3f:3d:21', 'NW-HK-SW1', 'USL48P', '192.168.1.52', true, now())"), {"g": gw})
        c.execute(text("INSERT INTO device_inventory (poll_device_id, mac, name, model, type, ip, connected, upgradable, seen_at) "
                       "VALUES (:g, 'aa:bb:cc:20:eb:bc', 'Door-1', 'UA-Hub', 'access', '192.168.1.53', false, true, now())"), {"g": gw})
    assert "NW-HK-SW1" not in (await client.get("/sources")).text            # Log sources: log devices only
    page = (await client.get("/devices")).text
    assert "NW-HK-SW1" in page and "Door-1" in page and "Access (door)" in page
    assert "/sources/add?register=192.168.1.52&amp;name=NW-HK-SW1&amp;model=USL48P" in page
    assert "register=192.168.1.53" not in page                                 # a door reader sends no syslog
    offline = (await client.get("/devices?status=offline")).text
    assert "Door-1" in offline and "NW-HK-SW1" not in offline
    assert "Door-1" not in (await client.get("/devices?q=usl48")).text
    assert "gw" in (await client.get("/devices?product=Log+source")).text     # registered, not reported by a controller
    prefilled = (await client.get("/sources/add?register=192.168.1.52&name=NW-HK-SW1&model=USL48P&vendor=unifi")).text
    assert 'value="NW-HK-SW1"' in prefilled and 'value="USL48P"' in prefilled


async def test_add_device_page_separate(client, make_user):
    await _admin(client, make_user)
    assert "Add device</button>" not in (await client.get("/sources")).text
    page = await client.get("/sources/add")
    assert page.status_code == 200 and "Add device</button>" in page.text
    assert 'class="side-item active" href="/sources/add"' in page.text


async def test_register_only_offered_where_device_can_send_itself(client, make_user, owner_engine):
    """Remote site: logs arrive from the gateway's public IP, so an AP's private IP is not offered for registration."""
    await _admin(client, make_user)
    with owner_engine.begin() as c:
        remote = c.execute(text("INSERT INTO clients (name, slug) VALUES ('CTS KL','cts-kl') RETURNING id")).scalar()
        local = c.execute(text("INSERT INTO clients (name, slug) VALUES ('Northwind HK','nwhk') RETURNING id")).scalar()
        gw_r = c.execute(text("INSERT INTO devices (client_id, name, vendor, syslog_source_ip) VALUES (:c,'CTS-KL','unifi','11.133.111.195') RETURNING id"),
                         {"c": remote}).scalar()
        gw_l = c.execute(text("INSERT INTO devices (client_id, name, vendor, syslog_source_ip) VALUES (:c,'NW-HK','unifi','192.168.1.1') RETURNING id"),
                         {"c": local}).scalar()
        for g, mac, name, ip in ((gw_r, "aa:00:00:a4:5a:12", "CTSKL-AP01", "192.168.2.61"), (gw_l, "aa:00:00:95:7c:fb", "NWHKWAP-04", "192.168.1.104")):
            c.execute(text("INSERT INTO device_inventory (poll_device_id, mac, name, model, type, ip, connected, seen_at) "
                           "VALUES (:g, :m, :n, 'U6 Lite', 'network', :ip, true, now())"), {"g": g, "m": mac, "n": name, "ip": ip})
    page = (await client.get("/devices")).text
    assert "register=192.168.1.104" in page                        # local site: the AP sends from its own address
    assert "register=192.168.2.61" not in page and "via CTS-KL" in page


async def test_all_devices_shows_live_cpu_memory_uptime(client, make_user, owner_engine):
    await _admin(client, make_user)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug) VALUES ('A','aa') RETURNING id")).scalar()
        gw = c.execute(text("INSERT INTO devices (client_id, name, vendor) VALUES (:c,'gw','unifi') RETURNING id"), {"c": cid}).scalar()
        c.execute(text("INSERT INTO device_inventory (poll_device_id, mac, name, model, type, ip, connected, seen_at, cpu_pct, mem_pct, uptime_s, client_count) "
                       "VALUES (:g, 'aa:bb:cc:37:79:5c', 'CTSKL-AP01', 'U6 Lite', 'uap', '192.168.2.61', true, now(), 12.3, 33.6, 2455105, 14)"), {"g": gw})
        c.execute(text("INSERT INTO device_inventory (poll_device_id, mac, name, model, type, ip, connected, seen_at, cpu_pct) "
                       "VALUES (:g, 'aa:bb:cc:c5:e0:52', 'stale-AP', 'U6 Lite', 'uap', '192.168.2.62', true, now() - interval '1 hour', 99)"), {"g": gw})
    page = (await client.get("/devices")).text
    row = page.split("CTSKL-AP01")[1].split("</tr>")[0]
    assert "12 %" in row and "34 %" in row and "28d 9h" in row and ">14<" in row
    assert "99 %" not in page.split("stale-AP")[1].split("</tr>")[0]           # values older than 5 minutes are not shown



async def test_asset_inventory_serial_lifecycle_and_excel(client, make_user, owner_engine):
    """Asset inventory: serial and UniFi's end-of-life flag shown, firmware-by-model compliance, Excel export."""
    import io

    from openpyxl import load_workbook
    from tests.conftest import full_login
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug) VALUES ('CTS KL','cts-kl') RETURNING id")).scalar()
        gw = c.execute(text("INSERT INTO devices (client_id, name, vendor) VALUES (:c, 'CTS-KL-UDM', 'unifi') RETURNING id"), {"c": cid}).scalar()
        for mac, name, fw, up, extra in (("60:22:32:16:06:64", "AP01", "6.7.57", False, '{"serial": "602232160664"}'),
                                          ("60:22:32:bb:d9:15", "AP02", "6.6.0", True, '{"serial": "602232BBD915", "eol": true}')):
            c.execute(text("""INSERT INTO device_inventory (poll_device_id, mac, name, model, type, firmware, upgradable, firmware_latest,
                                  connected, seen_at, extra) VALUES (:g, :m, :n, 'UAL6', 'uap', :fw, :up, '6.7.57', true, now(), CAST(:x AS jsonb))"""),
                      {"g": gw, "m": mac, "n": name, "fw": fw, "up": up, "x": extra})
    await full_login(client, make_user(role="engineer"))
    page = (await client.get("/devices")).text
    assert "602232BBD915" in page and ">EOL<" in page and "Firmware by model" in page and "End of life" in page
    assert "<code>6.6.0</code> ×1" in page                                       # two firmware versions for UAL6
    eol = (await client.get("/devices?lifecycle=eol")).text
    assert "AP02" in eol and "AP01" not in eol
    r = await client.get("/devices.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    rows = {row[1]: row for row in wb["Assets"].iter_rows(min_row=2, values_only=True)}
    assert rows["AP02"][4] == "602232BBD915" and rows["AP02"][8] == "6.7.57" and rows["AP02"][9] == "End of life"
    model = list(wb["Firmware by model"].iter_rows(min_row=2, values_only=True))[0]
    assert model[1] == "UAL6" and model[2] == 2 and model[4] == 1


async def test_site_overview(client, make_user, owner_engine):
    """One page per site: internet, VPN, devices (offline first), open alerts; other clients' sites are not reachable
    for a client user."""
    from tests.conftest import full_login
    with owner_engine.begin() as c:
        kl = c.execute(text("INSERT INTO clients (name, slug, timezone) VALUES ('CTS KL','cts-kl','Asia/Kuala_Lumpur') RETURNING id")).scalar()
        other = c.execute(text("INSERT INTO clients (name, slug) VALUES ('TSP-HK','tsp-hk') RETURNING id")).scalar()
        gw = c.execute(text("INSERT INTO devices (client_id, name, vendor, last_log_at) VALUES (:c, 'CTS-KL-UDM', 'unifi', now()) RETURNING id"), {"c": kl}).scalar()
        c.execute(text("INSERT INTO wan_state (client_id, device_id, wan_name, state, started_at) VALUES (:c, :d, 'internet', 'up', now() - interval '2 days')"), {"c": kl, "d": gw})
        c.execute(text("""INSERT INTO vpn_tunnel_info (device_id, tunnel_name, client_id, status, peer, updated_at)
                          VALUES (:d, 'ipsec:CTSKL-PartnerNet', :c, 'established', '11.229.44.205', now())"""), {"c": kl, "d": gw})
        c.execute(text("""INSERT INTO device_inventory (poll_device_id, mac, name, model, type, connected, seen_at, extra)
                          VALUES (:d, '60:22:32:16:06:64', 'CTSKL-AP01 - Workstation', 'UAL6', 'uap', false, now(),
                                  '{"uplink": {"device": "CTSKL-SW2", "port": 14}}')"""), {"d": gw})
        c.execute(text("""INSERT INTO alerts (rule_id, client_id, dedup_key, title, message, severity, status, first_seen, last_seen)
                          SELECT id, :c, 'k', 'Access point CTSKL-AP01 - Workstation is offline', 'm', 'high', 'open', now(), now()
                          FROM alert_rules LIMIT 1"""), {"c": kl})
    await full_login(client, make_user(role="engineer"))
    page = (await client.get(f"/site/{kl}")).text
    assert "CTS KL" in page and "CTSKL-PartnerNet" in page and "11.229.44.205" in page
    assert "CTSKL-AP01 - Workstation" in page and "CTSKL-SW2 port 14" in page and "▼ offline" in page
    assert "Access point CTSKL-AP01 - Workstation is offline" in page and "0 / 1" in page
    assert (await client.get("/site/999999")).status_code == 404
    token = csrf_from(page)                                                     # switching client on an overview opens the new one
    r = await client.post("/scope", data={"csrf_token": token, "client": str(other)}, headers={"referer": f"https://testserver/site/{kl}"})
    assert r.status_code == 303 and r.headers["location"] == f"/site/{other}"
    r = await client.post("/scope", data={"csrf_token": token, "client": "all"}, headers={"referer": f"https://testserver/site/{kl}"})
    assert r.headers["location"] == "/site"
    viewer = make_user(username="viewer-leo", role="client_readonly")
    with owner_engine.begin() as c:
        c.execute(text("INSERT INTO user_clients (user_id, client_id) VALUES (:u, :c)"), {"u": viewer["id"], "c": other})
    client.cookies.clear()
    await full_login(client, viewer)
    assert (await client.get(f"/site/{kl}")).status_code == 404

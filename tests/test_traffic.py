"""Internet traffic: driver counters, poller storage, report numbers vs. hand calculation, dashboard series."""

from datetime import UTC, datetime, timedelta

import pytest
from sqlalchemy import text

from app import db
from app.config import get_settings
from app.drivers.unifi import parse_wans
from app.poller.service import Poller
from app.reports.data import build, month_bounds
from tests.conftest import full_login
from tests.test_poller import FakeDriver, _cycle, site  # noqa: F401
from tests.test_wan import GATEWAY, NETWORKS, result

HK = "Asia/Hong_Kong"


def test_parse_wan_counters():
    g = dict(GATEWAY, wan2=dict(GATEWAY["wan2"], rx_bytes=10238580219977, tx_bytes=3344020305645, rx_packets=2381590223,
                                tx_packets=4419287941, rx_dropped=140079, tx_dropped=0, rx_errors=0, tx_errors=0, speed=1000))
    w = {x.name: x for x in parse_wans(g, NETWORKS)}["WAN2"]
    assert (w.rx_bytes, w.tx_packets, w.rx_dropped, w.link_speed_mbps) == (10238580219977, 4419287941, 140079, 1000)


async def test_poller_stores_counters(owner_engine, site):
    cid, gw, ap = site
    res = result()
    res.wans[1].rx_bytes, res.wans[1].tx_bytes, res.wans[1].rx_packets, res.wans[1].tx_packets = 1000, 500, 10, 5
    await _cycle(Poller(get_settings(), drivers={"fake": FakeDriver(res)}))
    with owner_engine.connect() as c:
        row = c.execute(text("SELECT wan_name, rx_bytes, tx_bytes, rx_packets FROM wan_metrics WHERE wan_name = 'WAN2'")).one()
    assert tuple(row) == ("WAN2", 1000, 500, 10)


@pytest.fixture
def october_traffic(owner_engine):
    """1 Oct 2026 (HKT) 10:00-11:00: 10 Mbps down / 2 Mbps up every minute, except one minute at 20 Mbps down.
    At 10:30 the gateway rebooted (counters restart near zero) — that minute must be skipped, not counted negative."""
    start, end = month_bounds(2026, 10, HK)
    with owner_engine.begin() as c:
        cid = c.execute(text("INSERT INTO clients (name, slug, timezone) VALUES ('Northwind HK', 'nwhk', :tz) RETURNING id"), {"tz": HK}).scalar()
        gw = c.execute(text("INSERT INTO devices (client_id, name, vendor) VALUES (:c, 'NW-HK', 'unifi') RETURNING id"), {"c": cid}).scalar()
        t0 = datetime(2026, 10, 1, 2, 0, tzinfo=UTC)          # 10:00 HKT
        rx = tx = prx = 0
        per_min_rx, per_min_tx = 75_000_000, 15_000_000     # 10 Mbps, 2 Mbps over 60 s
        for i in range(61):
            if i == 30:
                rx, tx, prx = 1_000, 1_000, 1                 # reboot: counters reset
            elif i > 0:
                rx += per_min_rx * (2 if i == 45 else 1)
                tx += per_min_tx
                prx += 1000
            c.execute(text("INSERT INTO wan_metrics (ts, client_id, device_id, wan_name, rx_bytes, tx_bytes, rx_packets, tx_packets, "
                           "rx_dropped, tx_dropped, rx_errors, tx_errors, link_speed_mbps) VALUES (:t, :c, :d, 'WAN2', :rx, :tx, :p, :p, 0, 0, 0, 0, 1000)"),
                      {"t": t0 + timedelta(minutes=i), "c": cid, "d": gw, "rx": rx, "tx": tx, "p": prx})
        c.execute(text("INSERT INTO wan_state (client_id, device_id, wan_name, label, state, started_at) VALUES (:c, :d, 'WAN2', 'Internet 2', 'up', :s)"),
                  {"c": cid, "d": gw, "s": start - timedelta(days=1)})
    return cid, start, end


async def test_internet_usage_matches_hand_calculation(october_traffic):
    cid, start, end = october_traffic
    async with db.sessionmaker()() as s:
        rep = await build(s, "internet_usage", cid, start, end, now=end + timedelta(days=1))
    (w,) = rep["sections"][0]["wans"]
    # 60 intervals, minute 30 skipped (reset) -> 59 counted; minute 45 carried double download.
    assert w["rx"] == 59 * 75_000_000 + 75_000_000 and w["tx"] == 59 * 15_000_000
    assert w["peak_rx_mbps"] == pytest.approx(20.0) and w["peak_tx_mbps"] == pytest.approx(2.0)
    assert w["p95_rx_mbps"] == pytest.approx(10.0) and w["avg_tx_mbps"] == pytest.approx(2.0)
    assert w["avg_rx_mbps"] == pytest.approx(60 * 75_000_000 * 8 / (59 * 60) / 1e6)
    assert w["label"] == "Internet 2" and w["availability"] == 1.0
    assert [(d["date"].isoformat(), d["rx"]) for d in w["days"]] == [("2026-10-01", 60 * 75_000_000)]


async def test_internet_usage_in_pdf_and_excel(october_traffic, owner_engine, tmp_path):
    cid, start, end = october_traffic
    from openpyxl import load_workbook
    from app.reporter.service import Reporter
    with owner_engine.begin() as c:
        c.execute(text("INSERT INTO report_runs (client_id, report_type, period_start, period_end) VALUES (:c, 'monthly_bundle', :s, :e)"),
                  {"c": cid, "s": start, "e": end})
    await Reporter(get_settings().model_copy(update={"reports_dir": str(tmp_path), "smtp_host": ""})).drain(now=end + timedelta(days=1))
    with owner_engine.connect() as c:
        run = c.execute(text("SELECT status, xlsx_path, summary, error FROM report_runs")).one()
    assert run.status == "done", run.error
    assert run.summary["internet_download_gb"] == 4.5 and abs(run.summary["internet_upload_gb"] - 0.885) <= 0.006
    wb = load_workbook(run.xlsx_path)
    row = list(wb["Internet usage"].iter_rows(min_row=2, values_only=True))[0]
    assert row[1] == "WAN2" and row[3] == 4.5 and row[9] == 20.0
    assert list(wb["Internet daily"].iter_rows(min_row=2, values_only=True))[0][3] == 4.5


async def test_dashboard_traffic_series(client, make_user, owner_engine, site):
    cid, gw, ap = site
    now = datetime.now(UTC).replace(second=0, microsecond=0)
    with owner_engine.begin() as c:
        for i in range(5):
            c.execute(text("INSERT INTO wan_metrics (ts, client_id, device_id, wan_name, rx_bytes, tx_bytes, rx_packets, tx_packets) "
                           "VALUES (:t, :c, :d, 'WAN2', :rx, :tx, :p, :p)"),
                      {"t": now - timedelta(minutes=10 - i), "c": cid, "d": gw, "rx": i * 75_000_000, "tx": i * 15_000_000, "p": i * 600})
    await full_login(client, make_user(role="engineer"))
    d = (await client.get("/api/dashboard?range=1h")).json()
    (s,) = d["traffic"]
    assert s["name"] == "WAN2" and s["rx_bytes"] == 4 * 75_000_000
    pts = [p for p in s["points"] if p[1] is not None]
    assert len(pts) == 4 and all(abs(p[1] - 10.0) < 1e-6 and abs(p[2] - 2.0) < 1e-6 and abs(p[3] - 10.0) < 1e-6 for p in pts)
    assert "Internet throughput" in (await client.get("/")).text
